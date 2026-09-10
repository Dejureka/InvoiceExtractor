"""GUI smoke tests — no display required."""

from __future__ import annotations

import importlib
import types
from pathlib import Path


def test_gui_module_imports():
    mod = importlib.import_module("invoice_extractor.gui")
    assert hasattr(mod, "run_gui")
    assert hasattr(mod, "_extract_one")
    assert hasattr(mod, "_extract_many")
    assert hasattr(mod, "_parse_dnd_paths")


def test_parse_dnd_paths_braces():
    from invoice_extractor.gui import _parse_dnd_paths

    assert _parse_dnd_paths(r"C:\a.pdf") == [r"C:\a.pdf"]
    assert _parse_dnd_paths(r"{C:\My Docs\inv.pdf}") == [r"C:\My Docs\inv.pdf"]
    parts = _parse_dnd_paths(r"{C:\a b.pdf} C:\c.pdf")
    assert parts == [r"C:\a b.pdf", r"C:\c.pdf"]


def test_cli_version_no_display():
    """CLI --version works without a display / GUI."""
    from invoice_extractor.cli import main

    try:
        code = main(["--version"])
    except SystemExit as e:
        assert e.code in (0, None)
        return
    assert code in (0, None)


def test_cli_extract_with_args_no_gui(tmp_path: Path, monkeypatch):
    """With PDF args, CLI extract path runs (no window)."""
    from invoice_extractor import cli as cli_mod

    sample = {
        "header": {"invoice_no": "X", "amount": 1.0},
        "items": [{"part_no": "P", "qty": 1}],
        "meta": {"format_id": "demo", "confidence": "rules", "source_file": "inv.pdf"},
    }

    class FakeResult:
        def to_dict(self):
            return dict(sample)

    pdf = tmp_path / "inv.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    out = tmp_path / "out.xlsx"

    monkeypatch.setattr(cli_mod, "extract_invoice", lambda *a, **k: FakeResult())
    monkeypatch.setattr(
        cli_mod, "hard_check", lambda data: {"verdict": "pass", "issues": []}
    )

    code = cli_mod.main([str(pdf), "--out", str(out)])
    assert code == 0
    assert out.is_file()


def test_cli_multi_pdfs_one_workbook(tmp_path: Path, monkeypatch):
    from invoice_extractor import cli as cli_mod
    from openpyxl import load_workbook
    from invoice_extractor.export import SUMMARY_SHEET

    def make_result(invoice_no: str):
        class FakeResult:
            def to_dict(self):
                return {
                    "header": {
                        "invoice_no": invoice_no,
                        "amount": 10.0,
                        "currency": "USD",
                        "item_line_count": 1,
                        "total_quantity": 1.0,
                    },
                    "items": [
                        {
                            "invoice_no": invoice_no,
                            "part_no": "P",
                            "qty": 1,
                            "amount": 10.0,
                        }
                    ],
                    "meta": {
                        "format_id": "demo",
                        "confidence": "rules",
                        "source_file": f"{invoice_no}.pdf",
                    },
                }

        return FakeResult()

    pdfs = []
    for name in ("a.pdf", "b.pdf"):
        p = tmp_path / name
        p.write_bytes(b"%PDF")
        pdfs.append(p)
    out = tmp_path / "batch.xlsx"

    def fake_extract(pdf, format_id=None):
        return make_result(Path(pdf).stem.upper())

    monkeypatch.setattr(cli_mod, "extract_invoice", fake_extract)
    monkeypatch.setattr(
        cli_mod, "hard_check", lambda data: {"verdict": "pass", "issues": []}
    )

    code = cli_mod.main([str(pdfs[0]), str(pdfs[1]), "--out", str(out)])
    assert code == 0
    assert out.is_file()
    wb = load_workbook(out)
    assert wb[SUMMARY_SHEET].max_row == 3


def test_extract_many_partial_failure(tmp_path: Path, monkeypatch):
    from invoice_extractor import gui as gui_mod
    from openpyxl import load_workbook
    from invoice_extractor.export import META_SHEET, SUMMARY_SHEET

    good = tmp_path / "good.pdf"
    bad = tmp_path / "bad.pdf"
    good.write_bytes(b"%PDF")
    bad.write_bytes(b"%PDF")
    out = tmp_path / "partial.xlsx"

    def fake_one(pdf: Path):
        if pdf.name == "bad.pdf":
            raise RuntimeError("parse failed")
        return {
            "header": {
                "invoice_no": "G1",
                "amount": 1.0,
                "currency": "USD",
                "item_line_count": 0,
                "total_quantity": 0,
            },
            "items": [],
            "meta": {"source_file": str(pdf), "format_id": "demo", "confidence": "rules"},
            "_hard": {"verdict": "pass", "issues": []},
        }

    monkeypatch.setattr(gui_mod, "_extract_one", fake_one)
    result = gui_mod._extract_many([good, bad], out)
    assert result["successes"] == 1
    assert len(result["failures"]) == 1
    assert Path(result["written"]).is_file()
    wb = load_workbook(result["written"])
    assert wb[SUMMARY_SHEET].max_row == 2
    assert wb[META_SHEET].max_row == 3


def _load_run_gui():
    import importlib.util
    import sys

    root = Path(__file__).resolve().parents[1]
    path = root / "run_gui.py"
    spec = importlib.util.spec_from_file_location("run_gui_entry", path)
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    sys.modules["run_gui_entry"] = mod
    spec.loader.exec_module(mod)
    return mod


def test_run_gui_routes_cli_when_args(monkeypatch):
    """run_gui.main with args delegates to CLI, does not open GUI."""
    rg = _load_run_gui()

    called: dict = {}

    def fake_cli(argv):
        called["argv"] = list(argv)
        return 7

    fake_mod = types.ModuleType("invoice_extractor.cli")
    fake_mod.main = fake_cli
    monkeypatch.setitem(__import__("sys").modules, "invoice_extractor.cli", fake_mod)

    def boom():
        raise AssertionError("GUI should not open when args present")

    fake_gui = types.ModuleType("invoice_extractor.gui")
    fake_gui.run_gui = boom
    monkeypatch.setitem(__import__("sys").modules, "invoice_extractor.gui", fake_gui)

    monkeypatch.setattr(
        __import__("sys"), "argv", ["run_gui.py", "doc.pdf", "--out", "o.xlsx"]
    )
    assert rg.main() == 7
    assert called["argv"] == ["doc.pdf", "--out", "o.xlsx"]


def test_run_gui_routes_gui_when_no_args(monkeypatch):
    """run_gui.main with no args opens GUI (mocked)."""
    rg = _load_run_gui()

    opened = {"ok": False}

    def fake_run_gui():
        opened["ok"] = True

    fake_gui = types.ModuleType("invoice_extractor.gui")
    fake_gui.run_gui = fake_run_gui
    monkeypatch.setitem(__import__("sys").modules, "invoice_extractor.gui", fake_gui)
    monkeypatch.setattr(__import__("sys"), "argv", ["run_gui.py"])
    assert rg.main() == 0
    assert opened["ok"] is True
