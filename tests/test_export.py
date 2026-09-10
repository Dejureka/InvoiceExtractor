"""Excel / JSON export + CLI default extension."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from invoice_extractor.cli import _out_path_for, build_parser
from invoice_extractor.export import default_out_path, write_extract, write_xlsx


SAMPLE = {
    "header": {
        "invoice_no": "INV-1",
        "invoice_date": "2026-01-02",
        "vendor": "Acme",
        "amount": 12.5,
        "currency": "USD",
        "total_quantity": 3.0,
        "item_line_count": 2,
        "total_pkg": 1.0,
        "gross_weight_kg": 2.5,
        "incoterm": "FOB",
        "origin": None,
        "hs_code": None,
    },
    "items": [
        {
            "invoice_no": "INV-1",
            "part_no": "A1",
            "description": "Widget",
            "qty": 2.0,
            "unit": "pcs",
            "unit_price": 5.0,
            "amount": 10.0,
            "currency": "USD",
            "origin": "CN",
            "hs_code": "1234",
        },
        {
            "invoice_no": "INV-1",
            "part_no": "B2",
            "description": "Gadget",
            "qty": 1.0,
            "unit": "pcs",
            "unit_price": 2.5,
            "amount": 2.5,
            "currency": "USD",
            "origin": "TW",
            "hs_code": "5678",
        },
    ],
    "meta": {
        "source_file": "inv.pdf",
        "format_id": "demo_v1",
        "text_backend": "fixture",
        "confidence": "high",
        "checker_verdict": "pass",
        "needs_gold": False,
        "needs_ocr": False,
        "notes": None,
        "checker_issues": [],
    },
}


def test_default_out_path_xlsx():
    assert default_out_path(Path("folder/inv.pdf")) == Path("folder/inv.extract.xlsx")
    assert _out_path_for(Path("a/b.pdf"), None) == Path("a/b.extract.xlsx")
    assert _out_path_for(Path("a/b.pdf"), "out.json") == Path("out.json")
    assert _out_path_for(Path("a/b.pdf"), "out.xlsx") == Path("out.xlsx")


def test_cli_help_mentions_xlsx():
    help_txt = build_parser().format_help()
    assert ".extract.xlsx" in help_txt or "xlsx" in help_txt.lower()


def test_write_xlsx_sheets(tmp_path: Path):
    path = tmp_path / "out.xlsx"
    write_xlsx(SAMPLE, path)
    assert path.is_file()
    wb = load_workbook(path)
    assert wb.sheetnames == ["header", "items", "meta"]

    ws_h = wb["header"]
    headers = [c.value for c in ws_h[1]]
    assert "invoice_no" in headers
    assert "amount" in headers
    assert "vendor" in headers
    row = {headers[i]: ws_h[2][i].value for i in range(len(headers))}
    assert row["invoice_no"] == "INV-1"
    assert row["amount"] == 12.5
    assert row["currency"] == "USD"

    ws_i = wb["items"]
    item_headers = [c.value for c in ws_i[1]]
    assert "part_no" in item_headers
    assert ws_i.max_row == 3  # header + 2 items
    part_idx = item_headers.index("part_no")
    assert ws_i[2][part_idx].value == "A1"
    assert ws_i[3][part_idx].value == "B2"

    ws_m = wb["meta"]
    meta = {ws_m[r][0].value: ws_m[r][1].value for r in range(2, ws_m.max_row + 1)}
    assert meta["format_id"] == "demo_v1"
    assert meta["checker_verdict"] == "pass"
    assert meta["source_file"] == "inv.pdf"


def test_write_extract_routes_by_extension(tmp_path: Path):
    xlsx = write_extract(SAMPLE, tmp_path / "a.xlsx")
    assert xlsx.suffix == ".xlsx"
    assert xlsx.is_file()
    wb = load_workbook(xlsx)
    assert "header" in wb.sheetnames

    jpath = write_extract(SAMPLE, tmp_path / "b.json")
    assert jpath.suffix == ".json"
    data = json.loads(jpath.read_text(encoding="utf-8"))
    assert data["header"]["invoice_no"] == "INV-1"

    # default non-json → xlsx (normalize .xls)
    xls = write_extract(SAMPLE, tmp_path / "c.xls")
    assert xls.suffix == ".xlsx"
    assert xls.is_file()
