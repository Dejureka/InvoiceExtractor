"""Convert Excel workbooks to layout-ish text for FormatSOP matching."""

from __future__ import annotations

from pathlib import Path

EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm", ".xls"})
OPENPYXL_SUFFIXES = frozenset({".xlsx", ".xlsm"})

# Tool outputs / templates that must not be treated as invoice inputs.
SKIP_INPUT_NAMES = frozenset(
    {
        "invoiceextract_result.xlsx",
        "invoiceextract_template.xlsx",
        "blextract_result.xlsx",
    }
)


class ExcelUnsupportedError(ValueError):
    """Raised for legacy .xls or unreadable workbooks with a user-facing message."""


def is_excel_path(path: str | Path) -> bool:
    return Path(path).suffix.lower() in EXCEL_SUFFIXES


def is_openpyxl_excel(path: str | Path) -> bool:
    return Path(path).suffix.lower() in OPENPYXL_SUFFIXES


def should_skip_input(path: str | Path) -> bool:
    """True for tool Result/Template Excel or anything under ``ocr_out/``."""
    p = Path(path)
    if p.name.lower() in SKIP_INPUT_NAMES:
        return True
    # Any path segment named ocr_out (portable or nested)
    return any(part.lower() == "ocr_out" for part in p.parts)


def workbook_to_text(path: str | Path) -> tuple[str, str]:
    """Return ``(text, backend)`` from an Excel workbook.

    Sheets become blocks headed by ``=== Sheet: <name> ===``; each row is
    non-empty cells joined by a tab. Empty rows are kept as blank lines so
    block structure stays visible for future Excel FormatSOP rules.
    """
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".xls":
        raise ExcelUnsupportedError(
            f"Legacy .xls not supported ({p.name}). Please save as .xlsx and retry."
        )
    if suffix not in OPENPYXL_SUFFIXES:
        raise ExcelUnsupportedError(f"Not an Excel workbook: {p}")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ExcelUnsupportedError(
            "openpyxl is required to read .xlsx/.xlsm inputs"
        ) from exc

    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelUnsupportedError(f"Cannot open Excel file {p.name}: {exc}") from exc

    blocks: list[str] = []
    try:
        for sheet in wb.worksheets:
            lines = [f"=== Sheet: {sheet.title} ==="]
            for row in sheet.iter_rows(values_only=True):
                cells: list[str] = []
                for cell in row:
                    if cell is None:
                        continue
                    s = str(cell).strip()
                    if s:
                        cells.append(s)
                lines.append("\t".join(cells) if cells else "")
            blocks.append("\n".join(lines).rstrip())
    finally:
        wb.close()

    text = "\n\n".join(blocks).strip() + ("\n" if blocks else "")
    return text, "excel/openpyxl"
